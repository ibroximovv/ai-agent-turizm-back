"""Document text extraction for PDF, PPTX, DOCX, TXT and MD."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

#: Below this, a PDF page is most likely a scanned image rather than text.
SPARSE_PAGE_THRESHOLD = 40

_MARKDOWN_HEADING_SPLIT = re.compile(r"(?m)(?=^#{1,3}\s)")
_BLANK_LINE_SPLIT = re.compile(r"(?:\r?\n){3,}")


class UnsupportedFormatError(ValueError):
    """Raised for an extension the pipeline has no parser for."""


@dataclass(frozen=True)
class ParsedChunk:
    #: e.g. "page 1", "slide 3", "section 1"
    label: str
    text: str


@dataclass
class ParsedDocument:
    chunks: list[ParsedChunk] = field(default_factory=list)
    total_chars: int = 0
    warnings: list[str] = field(default_factory=list)


def parse_file(file_path: Path | str) -> ParsedDocument:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in (".txt", ".md"):
        return _parse_text_file(path)

    raise UnsupportedFormatError(
        f"Qo'llab-quvvatlanmaydigan fayl formati: {suffix}. "
        "Faqat .pdf, .pptx, .docx, .txt va .md qabul qilinadi."
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    warnings: list[str] = []
    chunks: list[ParsedChunk] = []
    total_chars = 0

    try:
        reader = PdfReader(str(path))
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if not page_text:
                continue

            if len(page_text) < SPARSE_PAGE_THRESHOLD:
                warnings.append(
                    f"{index}-sahifada juda kam matn bor ({len(page_text)} belgi) — "
                    "skanerlangan rasm bo'lishi mumkin."
                )

            chunks.append(ParsedChunk(label=f"page {index}", text=page_text))
            total_chars += len(page_text)
    except Exception as exc:
        logger.error("PDF parse xatosi %s: %s", path, exc)
        raise RuntimeError(f"PDF o'qib bo'lmadi: {exc}") from exc

    if not chunks:
        warnings.append(
            "PDF ichidan matn topilmadi — hujjat skanerlangan va OCR talab qiladi."
        )

    return ParsedDocument(chunks=chunks, total_chars=total_chars, warnings=warnings)


def _parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    warnings: list[str] = []
    chunks: list[ParsedChunk] = []
    total_chars = 0

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    if not slides:
        warnings.append("Taqdimotda slaydlar topilmadi.")

    for index, slide in enumerate(slides, start=1):
        slide_text = "\n".join(_shape_paragraphs(slide.shapes))

        notes_text = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes_text = " ".join(notes_frame.text.split())

        combined = slide_text
        if notes_text:
            combined += f"\n\n[Presenter Notes: {notes_text}]"
        combined = combined.strip()

        if combined:
            chunks.append(ParsedChunk(label=f"slide {index}", text=combined))
            total_chars += len(combined)

    return ParsedDocument(chunks=chunks, total_chars=total_chars, warnings=warnings)


def _shape_paragraphs(shapes) -> list[str]:
    """One entry per paragraph, preserving each shape's line structure.

    Group shapes are walked recursively; tables contribute one line per row.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_paragraphs(shape.shapes))
            continue

        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    lines.append(row_text)
            continue

        if not getattr(shape, "has_text_frame", False):
            continue

        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                lines.append(text)

    return lines


def _parse_docx(path: Path) -> ParsedDocument:
    import mammoth

    try:
        with path.open("rb") as handle:
            result = mammoth.convert_to_markdown(handle)
    except Exception as exc:
        raise RuntimeError(f"DOCX o'qib bo'lmadi: {exc}") from exc

    markdown = result.value or ""
    warnings = [message.message for message in result.messages if getattr(message, "message", None)]

    # Split into sections at markdown headings.
    chunks = _to_sequential_chunks(_MARKDOWN_HEADING_SPLIT.split(markdown), "section")
    if not chunks and markdown.strip():
        chunks = [ParsedChunk(label="section 1", text=markdown.strip())]

    return ParsedDocument(
        chunks=chunks,
        total_chars=sum(len(chunk.text) for chunk in chunks),
        warnings=warnings,
    )


def _parse_text_file(path: Path) -> ParsedDocument:
    content = path.read_text(encoding="utf-8", errors="replace")
    chunks = _to_sequential_chunks(_BLANK_LINE_SPLIT.split(content), "part")
    if not chunks and content.strip():
        chunks = [ParsedChunk(label="part 1", text=content.strip())]

    return ParsedDocument(
        chunks=chunks,
        total_chars=sum(len(chunk.text) for chunk in chunks),
        warnings=[],
    )


def _to_sequential_chunks(raw_sections: list[str], label_prefix: str) -> list[ParsedChunk]:
    """Number the labels over the kept sections only.

    Otherwise the emitted grounding markers would reference a section that is
    missing from the document.
    """
    chunks: list[ParsedChunk] = []
    for raw in raw_sections:
        text = raw.strip()
        if not text:
            continue
        chunks.append(ParsedChunk(label=f"{label_prefix} {len(chunks) + 1}", text=text))
    return chunks
